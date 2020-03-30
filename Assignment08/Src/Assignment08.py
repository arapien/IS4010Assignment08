'''
Andrea Rapien
rapienac@mail.uc.edu
Assignment 08
IS 4010
Spring 2020
This project demonstrates the ability to modify a PyDev Project to process an existing spreadsheet, generate a chart, and save it in a new excel sheet.
@author: rapienac
'''
'''
Created on Mar 20, 2020
Assignment 08 Donor Code
Adapted from https://openpyxl.readthedocs.io/en/stable/charts/pie.html
@author: nicomp
'''
from openpyxl import load_workbook
from openpyxl.chart import (
    PieChart,
    ProjectedPieChart,
    Reference
)
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList 

wb = load_workbook(filename = 'Top5TransactionsByLoyaltyNumber.xlsx')
ws = wb['Sheet1']

pie = PieChart()
labels = Reference(ws, min_col=4, min_row=2, max_row=6)
data = Reference(ws, min_col=2, min_row=2, max_row=6)
pie.add_data(data, titles_from_data=False)
pie.set_categories(labels)
pie.title = "Top 5 Total Transactions by Store"
pie.dataLabels = DataLabelList()
pie.dataLabels.showVal = True

# Cut the first slice out of the pie
pieSlice = DataPoint(idx=0, explosion=20)
pie.series[0].data_points = [pieSlice]

ws.add_chart(pie, "A8")
wb.save('Top5TransactionsByLoyaltyNumberWithPieChart.xlsx') # .xlsx file cannot be open when we do this
print("Done")
